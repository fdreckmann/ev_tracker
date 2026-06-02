"""
Tests for the central report-Excel helper and the automatic e-mail report's
use of it (Excel template attachments).
"""
import io
import json

import openpyxl


def _seed_session(app, year=2026, month=5, n=1):
    from core.db import _get_db, close_db_if_owned
    with app.app_context():
        con = _get_db()
        for i in range(n):
            start = f"{year:04d}-{month:02d}-{1+i:02d}T10:00:00"
            end   = f"{year:04d}-{month:02d}-{1+i:02d}T11:00:00"
            con.execute("""INSERT INTO sessions
                (start_ts, end_ts, kwh_charged, cost_eur, cost_manual,
                 location, charger_type, vehicle_id, provider)
                VALUES (?,?,?,?,?,?,?,?,?)""",
                (start, end, 20.5, 6.15, 1, "home", "ac", "v0", "manual"))
        con.commit()
        close_db_if_owned(con)


def _write_template(app, sheet="Tabelle1"):
    """Write a minimal template.xlsx where export_excel expects it."""
    import export_excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws["A1"] = "Ladeprotokoll"
    ws["A5"] = "Datum"; ws["B5"] = "kWh"; ws["C5"] = "Kosten"
    wb.save(export_excel.TEMPLATE_PATH)
    return str(export_excel.TEMPLATE_PATH)


def _set_cfg(app, **kv):
    from core.config import load_config, save_config
    with app.app_context():
        cfg = load_config()
        cfg.update(kv)
        save_config(cfg)


# ── 1. Helper produces a valid XLSX ───────────────────────────────────────────
class TestBuildHelper:
    def test_builtin_returns_valid_xlsx(self, app):
        _seed_session(app)
        with app.app_context():
            from core.config import load_config
            from services.report_excel_service import build_report_excel_bytes
            from datetime import date
            cfg = load_config()
            xb, warns = build_report_excel_bytes(
                {"start": date(2026, 5, 1)}, "all", "all", cfg, "de",
                template_id="builtin:standard")
        wb = openpyxl.load_workbook(io.BytesIO(xb))
        assert len(wb.sheetnames) >= 1
        assert isinstance(warns, list)

    def test_active_template_used_when_present(self, app):
        _seed_session(app)
        _write_template(app)
        _set_cfg(app,
                 template_column_mapping={"1": "datum", "2": "kwh", "3": "kosten"},
                 template_start_row=6, template_header_row=5,
                 template_sheet="Tabelle1")
        with app.app_context():
            from core.config import load_config
            from services.report_excel_service import (
                build_report_excel_bytes, resolve_template_settings)
            from datetime import date
            cfg = load_config()
            settings = resolve_template_settings(cfg, "active")
            assert settings["kind"] == "active"
            assert settings["col_override"] == {"1": "datum", "2": "kwh", "3": "kosten"}
            assert settings["start_row"] == 6
            assert settings["sheet"] == "Tabelle1"
            xb, _w = build_report_excel_bytes(
                {"start": date(2026, 5, 1)}, "all", "all", cfg, "de",
                template_id="active")
        wb = openpyxl.load_workbook(io.BytesIO(xb))
        assert "Tabelle1" in wb.sheetnames  # template sheet preserved


# ── 2/3. Auto single-month report attaches XLSX built from the helper ─────────
class TestAutoReportExcel:
    def _smtp_cfg(self):
        return dict(smtp_host="localhost", smtp_from_email="ev@test.local",
                    report_email_enabled=True,
                    report_email_recipients=["a@test.local"])

    def test_single_month_attaches_xlsx(self, app, monkeypatch):
        _seed_session(app)
        _set_cfg(app, **self._smtp_cfg())
        captured = {}
        import routes.email_reports as er

        def fake_send(to, subject, html, attachments):
            captured["attachments"] = attachments
            return True, None
        monkeypatch.setattr(er, "_send_email_with_attachments", fake_send)
        with app.app_context():
            from core.config import load_config
            cfg = load_config()
            cfg["report_email_period_mode"] = "single_month"
            cfg["report_email_single_month"] = "2026-05"
            cfg["report_email_include_excel"] = True
            ok, err = er._send_report_email(cfg=cfg, triggered_by="manual")
        assert ok, err
        atts = captured.get("attachments", [])
        assert any(a[0].endswith(".xlsx") for a in atts), "XLSX attachment missing"
        xb = next(a[1] for a in atts if a[0].endswith(".xlsx"))
        openpyxl.load_workbook(io.BytesIO(xb))  # must be a valid workbook

    def test_no_typeerror_from_config_kwarg(self, app, monkeypatch):
        """Regression: export() must never be called with config=cfg."""
        _seed_session(app)
        _set_cfg(app, **self._smtp_cfg())
        import routes.email_reports as er
        monkeypatch.setattr(er, "_send_email_with_attachments",
                            lambda *a, **k: (True, None))
        # Wrap export() to fail loudly if a 'config' kwarg ever sneaks back in.
        import export_excel
        orig = export_excel.export
        def guarded(*a, **k):
            assert "config" not in k, "export() called with forbidden config kwarg"
            return orig(*a, **k)
        monkeypatch.setattr(export_excel, "export", guarded)
        with app.app_context():
            from core.config import load_config
            cfg = load_config()
            cfg["report_email_period_mode"] = "single_month"
            cfg["report_email_single_month"] = "2026-05"
            ok, err = er._send_report_email(cfg=cfg, triggered_by="manual")
        assert ok, err

    def test_uses_saved_template_mapping(self, app, monkeypatch):
        _seed_session(app)
        _write_template(app)
        _set_cfg(app,
                 template_column_mapping={"1": "datum", "2": "kwh"},
                 template_start_row=6, template_header_row=5,
                 template_cell_mapping={"A1": "kennzeichen"},
                 template_sheet="Tabelle1",
                 template_footer_start_row=20,
                 # hash agreement so mapping is considered valid
                 active_template={"source": "upload", "name": "T", "hash": "abc"},
                 template_mapping_hash="abc",
                 **self._smtp_cfg())
        seen = {}
        import export_excel
        orig = export_excel.export
        def spy(*a, **k):
            seen.update(k)
            return orig(*a, **k)
        monkeypatch.setattr(export_excel, "export", spy)
        import routes.email_reports as er
        monkeypatch.setattr(er, "_send_email_with_attachments",
                            lambda *a, **k: (True, None))
        with app.app_context():
            from core.config import load_config
            cfg = load_config()
            cfg["report_email_period_mode"] = "single_month"
            cfg["report_email_single_month"] = "2026-05"
            ok, err = er._send_report_email(cfg=cfg, triggered_by="manual")
        assert ok, err
        assert seen.get("col_override") == {"1": "datum", "2": "kwh"}
        assert seen.get("start_row") == 6
        assert seen.get("header_row") == 5
        assert seen.get("cell_mapping") == {"A1": "kennzeichen"}
        assert seen.get("sheet") == "Tabelle1"
        assert seen.get("footer_start_row") == 20


# ── 4. Hash mismatch / broken mapping → report fails, no silent success ───────
class TestValidationBlocksBrokenAttachment:
    def test_hash_mismatch_fails_report(self, app, monkeypatch):
        _seed_session(app)
        _write_template(app)
        _set_cfg(app,
                 template_column_mapping={"1": "datum"},
                 template_start_row=6,
                 active_template={"source": "upload", "name": "T", "hash": "NEWHASH"},
                 template_mapping_hash="OLDHASH",
                 smtp_host="localhost", smtp_from_email="ev@test.local",
                 report_email_enabled=True,
                 report_email_recipients=["a@test.local"],
                 report_email_period_mode="single_month",
                 report_email_single_month="2026-05",
                 report_email_include_excel=True)
        import routes.email_reports as er
        sent = {"called": False}
        def fake_send(*a, **k):
            sent["called"] = True
            return True, None
        monkeypatch.setattr(er, "_send_email_with_attachments", fake_send)
        with app.app_context():
            from core.config import load_config
            ok, err = er._send_report_email(cfg=load_config(), triggered_by="manual")
        assert ok is False
        assert "Mapping" in (err or "")
        assert sent["called"] is False  # never sent a broken attachment

    def test_history_records_error(self, app, monkeypatch):
        _seed_session(app)
        _write_template(app)
        _set_cfg(app,
                 template_column_mapping={"1": "datum"},
                 template_start_row=6,
                 active_template={"source": "upload", "name": "T", "hash": "NEWHASH"},
                 template_mapping_hash="OLDHASH",
                 smtp_host="localhost", smtp_from_email="ev@test.local",
                 report_email_enabled=True,
                 report_email_recipients=["a@test.local"],
                 report_email_period_mode="single_month",
                 report_email_single_month="2026-05",
                 report_email_include_excel=True)
        import routes.email_reports as er
        monkeypatch.setattr(er, "_send_email_with_attachments",
                            lambda *a, **k: (True, None))
        with app.app_context():
            from core.config import load_config
            from core.db import _get_db, close_db_if_owned
            er._send_report_email(cfg=load_config(), triggered_by="manual")
            con = _get_db()
            row = con.execute(
                "SELECT status, excel_error FROM email_report_history "
                "ORDER BY id DESC LIMIT 1").fetchone()
            close_db_if_owned(con)
        assert row["status"] == "error"
        assert row["excel_error"]


# ── 5. Archive report and auto report share the helper ───────────────────────
class TestSharedHelper:
    def test_archive_and_auto_share_helper(self, app, authed_client, monkeypatch):
        """Both /api/reports/create and the auto-report call the same helper
        (routes import it lazily, so patching the module attribute is enough)."""
        _seed_session(app)
        import services.report_excel_service as svc
        orig = svc.build_report_excel_bytes
        calls = {"n": 0}
        def counted(*a, **k):
            calls["n"] += 1
            return orig(*a, **k)
        monkeypatch.setattr(svc, "build_report_excel_bytes", counted)

        # archive report (single month)
        rv = authed_client.post("/api/reports/create",
                                json={"year": 2026, "month": 5, "include_excel": True})
        assert rv.status_code == 200 and rv.get_json()["has_excel"] is True
        archive_calls = calls["n"]
        assert archive_calls >= 1

        # auto report (single month)
        import routes.email_reports as er
        monkeypatch.setattr(er, "_send_email_with_attachments",
                            lambda *a, **k: (True, None))
        with app.app_context():
            from core.config import load_config
            cfg = load_config()
            cfg.update(smtp_host="localhost", smtp_from_email="ev@test.local",
                       report_email_enabled=True,
                       report_email_recipients=["a@test.local"],
                       report_email_period_mode="single_month",
                       report_email_single_month="2026-05",
                       report_email_include_excel=True)
            ok, err = er._send_report_email(cfg=cfg, triggered_by="manual")
        assert ok, err
        assert calls["n"] > archive_calls  # auto report also used the helper

    def test_export_route_no_config_kwarg(self, authed_client, app):
        _seed_session(app)
        rv = authed_client.get("/api/export?year=2026&month=5&location=all")
        assert rv.status_code == 200
        assert rv.mimetype == ("application/vnd.openxmlformats-"
                               "officedocument.spreadsheetml.sheet")


# ── 6. Config round-trips report_email_template_id ───────────────────────────
class TestConfigRoundtrip:
    def test_template_id_saved_and_loaded(self, authed_client, app):
        rv = authed_client.post("/api/report/config",
                                json={"report_email_template_id": "builtin:standard"})
        assert rv.status_code == 200
        got = authed_client.get("/api/report/config").get_json()
        assert got["report_email_template_id"] == "builtin:standard"

    def test_multi_month_mode_saved(self, authed_client, app):
        authed_client.post("/api/report/config",
                           json={"report_email_multi_month_excel_mode": "template_per_month_zip"})
        got = authed_client.get("/api/report/config").get_json()
        assert got["report_email_multi_month_excel_mode"] == "template_per_month_zip"

    def test_template_status_route(self, authed_client, app):
        rv = authed_client.get("/api/report/template-status")
        assert rv.status_code == 200
        data = rv.get_json()
        assert "status" in data and "template_name" in data and "templates" in data


# ── 7. UI contains a template selector + preview/test buttons ────────────────
class TestReportUI:
    def _read(self, *parts):
        from pathlib import Path
        return (Path(__file__).parent.parent / "app").joinpath(*parts).read_text()

    def test_index_has_template_selector(self):
        html = self._read("templates", "index.html")
        assert 'id="rep_template_id"' in html
        assert 'id="rep_template_status"' in html

    def test_index_has_preview_and_testmail_buttons(self):
        html = self._read("templates", "index.html")
        assert "downloadReportExcelPreview()" in html
        assert "sendReportTestMail()" in html

    def test_reports_js_functions_present(self):
        js = self._read("static", "js", "reports.js")
        for fn in ("loadReportTemplateStatus", "sendReportTestMail",
                   "downloadReportExcelPreview", "onRepExcelToggle"):
            assert f"function {fn}" in js, fn

    def test_reports_js_saves_template_id(self):
        js = self._read("static", "js", "reports.js")
        assert "report_email_template_id" in js
        assert "report_email_multi_month_excel_mode" in js


# ── 8/9. Multi-month modes ───────────────────────────────────────────────────
class TestMultiMonth:
    def test_standard_multi_sheet(self, app, monkeypatch):
        _seed_session(app, month=4)
        _seed_session(app, month=5)
        import routes.email_reports as er
        captured = {}
        monkeypatch.setattr(er, "_send_email_with_attachments",
                            lambda to, s, h, atts: (captured.update(a=atts) or (True, None)))
        with app.app_context():
            from core.config import load_config
            cfg = load_config()
            cfg.update(smtp_host="localhost", smtp_from_email="ev@test.local",
                       report_email_enabled=True,
                       report_email_recipients=["a@test.local"],
                       report_email_period_mode="multiple_months",
                       report_email_months=["2026-04", "2026-05"],
                       report_email_include_excel=True,
                       report_email_multi_month_excel_mode="standard_multi_sheet")
            ok, err = er._send_report_email(cfg=cfg, triggered_by="manual")
        assert ok, err
        atts = captured.get("a", [])
        assert any(a[0].endswith(".xlsx") for a in atts)

    def test_template_per_month_zip(self, app, monkeypatch):
        _seed_session(app, month=4)
        _seed_session(app, month=5)
        import routes.email_reports as er
        captured = {}
        monkeypatch.setattr(er, "_send_email_with_attachments",
                            lambda to, s, h, atts: (captured.update(a=atts) or (True, None)))
        with app.app_context():
            from core.config import load_config
            cfg = load_config()
            cfg.update(smtp_host="localhost", smtp_from_email="ev@test.local",
                       report_email_enabled=True,
                       report_email_recipients=["a@test.local"],
                       report_email_period_mode="multiple_months",
                       report_email_months=["2026-04", "2026-05"],
                       report_email_include_excel=True,
                       report_email_multi_month_excel_mode="template_per_month_zip")
            ok, err = er._send_report_email(cfg=cfg, triggered_by="manual")
        assert ok, err
        atts = captured.get("a", [])
        zips = [a for a in atts if a[0].endswith(".zip")]
        assert zips, "expected a ZIP attachment"
        import zipfile
        zf = zipfile.ZipFile(io.BytesIO(zips[0][1]))
        xlsx_members = [n for n in zf.namelist() if n.endswith(".xlsx")]
        assert len(xlsx_members) == 2  # one per month
